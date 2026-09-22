#!/usr/bin/env python3
"""
Extractor de la hoja INVENTARIO del libro de inventario a `inventario_apertura.csv`.

POR QUÉ EXISTE.
Hasta ahora `inventario_apertura.csv` aparecía en el repositorio sin decir cómo
se había producido. Para GAMP 5 eso es un hueco: el registro de
`gmp.migracion_apertura` guarda el hash del archivo de origen como evidencia de
integridad, y esa evidencia no vale nada si el paso del .xlsx al .csv es un
recorte manual que nadie puede repetir. Este script cierra la cadena:

    05 INVENTARIO NAIL SHOW FABRICA.xlsx
      -> xlsx_a_csv.py     -> inventario_apertura.csv + origen.json
      -> generar_migracion.mjs -> supabase/migrations/<ts>_carga_saldo_apertura.sql
      -> supabase db push

`origen.json` lleva el nombre y el sha256 del **xlsx**, no del csv: el xlsx es
la fuente, el csv es un intermedio de esta herramienta.

DEPENDENCIA.
Usa `openpyxl`, que NO es dependencia del proyecto (CLAUDE.md §7) sino de esta
herramienta de línea de comandos, que corre a mano y fuera del build y del
bundle. Instalar con `pip install openpyxl` antes de usarlo.

QUÉ COLUMNA ES QUÉ.
La hoja tiene un encabezado en dos niveles: la fila 1 abre `CONTENIDO` sobre
las columnas D y E, y la fila 2 las nombra `PESO` y `CANTIDAD`. Los datos
arrancan en la fila 3.

    A CODIGO   B INSUMO(familia)  C PROVEEDOR   D PESO   E CANTIDAD
    F DETALLE  G FECHA ACT        H COSTO U$D   I COSTO $   J Unitario

`CANTIDAD` se toma como las unidades en existencia, por indicación de la
conducción del proyecto (2026-09-22). Ver la nota de la migración generada: la
lectura no es evidente en la planilla y quedó asentada como suposición.

El color de fondo de la fila se conserva crudo en `color_origen` porque nadie
definió todavía qué codifica (D-22 de docs/DECISIONES_ABIERTAS.md). Se guarda
sin interpretar para no inventar la regla.

Uso:
    python scripts/apertura/xlsx_a_csv.py "<ruta al .xlsx>"
"""

import csv
import hashlib
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - mensaje para quien corre el script
    sys.exit("Falta openpyxl. Instalalo con: pip install openpyxl")

AQUI = Path(__file__).resolve().parent
CSV_SALIDA = AQUI / "inventario_apertura.csv"
ORIGEN_SALIDA = AQUI / "origen.json"

HOJA = "INVENTARIO"
PRIMERA_FILA_DATOS = 3

ENCABEZADO = [
    "codigo",
    "familia",
    "proveedor",
    "contenido_peso",
    "contenido_cantidad",
    "detalle",
    "fecha_actualizacion_costo",
    "costo_usd",
    "costo_ars",
    "costo_unitario_ars",
    "color_origen",
]

# Colores literales que usa la planilla. El nombre es descriptivo del color, no
# de un significado: qué codifica cada uno sigue siendo D-22.
COLORES_RGB = {
    "FF00B050": "VERDE",
    "FFFFFF00": "AMARILLO",
    "FFFF0000": "ROJO",
    "FF7030A0": "VIOLETA",
    "FF00B0F0": "CELESTE",
}


def color_de_fila(celda) -> str:
    """Nombre del color de fondo, o SIN_COLOR / TEMA_n para lo que no se mapea."""
    relleno = celda.fill
    if relleno is None or relleno.patternType is None:
        return "SIN_COLOR"

    rgb = getattr(relleno.fgColor, "rgb", None)
    if isinstance(rgb, str):
        if rgb in COLORES_RGB:
            return COLORES_RGB[rgb]
        if rgb in ("00000000", "FFFFFFFF"):
            return "SIN_COLOR"
        return f"RGB_{rgb}"

    tema = getattr(relleno.fgColor, "theme", None)
    if isinstance(tema, int):
        tinte = getattr(relleno.fgColor, "tint", 0.0) or 0.0
        # El gris de la planilla es el tema 0 oscurecido. Aparece en la versión
        # (4) y no existía en la anterior.
        if tema == 0 and tinte < 0:
            return "GRIS"
        return f"TEMA_{tema}"

    return "SIN_COLOR"


def texto(valor) -> str:
    if valor is None:
        return ""
    if hasattr(valor, "date"):  # datetime de la columna FECHA ACT
        return valor.date().isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(f"Uso: python {Path(__file__).name} \"<ruta al .xlsx>\"")

    ruta = Path(sys.argv[1])
    if not ruta.is_file():
        sys.exit(f"No existe el archivo: {ruta}")

    sha256 = hashlib.sha256(ruta.read_bytes()).hexdigest()

    # data_only=True toma el valor calculado de las fórmulas, no la fórmula.
    # Sin read_only, porque hace falta leer el relleno de cada celda.
    libro = openpyxl.load_workbook(ruta, data_only=True)
    if HOJA not in libro.sheetnames:
        sys.exit(f"El libro no tiene una hoja {HOJA!r}. Tiene: {libro.sheetnames}")
    hoja = libro[HOJA]

    filas = []
    vacias = 0
    for i in range(PRIMERA_FILA_DATOS, hoja.max_row + 1):
        valores = [hoja.cell(row=i, column=c).value for c in range(1, 11)]
        if valores[0] is None or not str(valores[0]).strip():
            vacias += 1
            continue
        filas.append([texto(v) for v in valores] + [color_de_fila(hoja.cell(row=i, column=1))])

    with CSV_SALIDA.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(ENCABEZADO)
        w.writerows(filas)

    ORIGEN_SALIDA.write_text(
        json.dumps(
            {
                "archivo_origen": ruta.name,
                "hash_archivo": sha256,
                "hoja": HOJA,
                "renglones": len(filas),
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )

    codigos = {f[0] for f in filas}
    print(f"  origen      : {ruta.name}")
    print(f"  sha256      : {sha256}")
    print(f"  renglones   : {len(filas)}  (filas sin código salteadas: {vacias})")
    print(f"  códigos     : {len(codigos)} distintos")
    print(f"  escrito     : {CSV_SALIDA.name}, {ORIGEN_SALIDA.name}")


if __name__ == "__main__":
    main()
